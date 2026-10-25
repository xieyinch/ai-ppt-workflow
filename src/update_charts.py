"""Update native chart data in a generated PPTX.
Input content JSON uses slide source numbers and chart occurrence/name.
"""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
except ModuleNotFoundError:
    sys.path.insert(0, str(ROOT / ".deps"))
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData


def update(pptx_path: Path, content_path: Path, source_indices: bool = False) -> None:
    prs = Presentation(str(pptx_path))
    cfg = json.loads(content_path.read_text(encoding="utf-8"))
    for output_index, slide_cfg in enumerate(cfg.get("slides", [])):
        entries = slide_cfg.get("charts", []) or []
        if not entries:
            continue
        slide_index = int(slide_cfg["src"]) - 1 if source_indices else output_index
        if slide_index < 0 or slide_index >= len(prs.slides):
            raise ValueError(f"chart slide out of range: {slide_index + 1}")
        slide = prs.slides[slide_index]
        charts = [shape for shape in slide.shapes if getattr(shape, "has_chart", False)]
        for entry in entries:
            target = None
            if entry.get("name"):
                target = next((s for s in charts if s.name == entry["name"]), None)
            else:
                occ = int(entry.get("occurrence", 1)) - 1
                if 0 <= occ < len(charts):
                    target = charts[occ]
            if target is None:
                raise ValueError(f"chart target not found on slide {slide_cfg['src']}")
            data = CategoryChartData()
            categories = list(entry.get("categories", []))
            data.categories = categories
            series = entry.get("series", {})
            for name, values in series.items():
                values = list(values)
                if len(values) != len(categories):
                    raise ValueError(f"series length mismatch: {name}")
                data.add_series(str(name), values)
            target.chart.replace_data(data)
    prs.save(str(pptx_path))


if __name__ == "__main__":
    if len(sys.argv) not in (3, 4):
        raise SystemExit("usage: update_charts.py OUTPUT.pptx CONTENT.json [--source]")
    update(Path(sys.argv[1]), Path(sys.argv[2]), len(sys.argv) == 4 and sys.argv[3] == "--source")
